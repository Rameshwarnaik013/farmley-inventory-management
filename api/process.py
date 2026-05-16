from fastapi import FastAPI, UploadFile, File, Form
from fastapi.responses import JSONResponse, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from io import BytesIO
from math import erf, sqrt
import json

app = FastAPI()
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])

CACHE = {}


def norm_cdf(z):
    return 0.5 * (1 + erf(z / sqrt(2)))


def compute_inventory(sigma_df, abc_df, idr_df, lead_time, z_scores, order_cycle):
    months = [c for c in sigma_df.columns if c not in ['Item_Name', 'New MIS ITEM Group', 'Metric']]
    qty_df = sigma_df[sigma_df['Metric'] == 'Qty (UNT)'].copy().reset_index(drop=True)

    qty_df['Avg_Monthly_Demand'] = qty_df[months].mean(axis=1)
    qty_df['Std_Dev_Demand'] = qty_df[months].std(axis=1)
    qty_df['CoV'] = qty_df['Std_Dev_Demand'] / qty_df['Avg_Monthly_Demand'].replace(0, np.nan)
    qty_df['Max_Demand'] = qty_df[months].max(axis=1)
    qty_df['Min_Demand'] = qty_df[months].min(axis=1)

    recent = months[-6:] if len(months) >= 6 else months
    qty_df['Recent_Avg'] = qty_df[recent].mean(axis=1)
    qty_df['Recent_Std'] = qty_df[recent].std(axis=1)

    abc_merge = abc_df[['Item_Name', 'Category ', 'Origin', 'Mrp', 'Usage']].copy()
    abc_merge.columns = ['Item_Name', 'ABC_Category', 'Origin', 'MRP', 'Annual_Usage_Value']

    inv = qty_df[['Item_Name', 'New MIS ITEM Group'] + months +
                 ['Avg_Monthly_Demand', 'Std_Dev_Demand', 'CoV', 'Max_Demand', 'Min_Demand',
                  'Recent_Avg', 'Recent_Std']].merge(abc_merge, on='Item_Name', how='left')

    z_map = {'A': z_scores[0], 'B': z_scores[1], 'C': z_scores[2]}
    inv['Z_Score'] = inv['ABC_Category'].map(z_map).fillna(z_scores[1])
    inv['Service_Level'] = inv['ABC_Category'].map({
        'A': f"{norm_cdf(z_scores[0]):.0%}",
        'B': f"{norm_cdf(z_scores[1]):.0%}",
        'C': f"{norm_cdf(z_scores[2]):.0%}"
    }).fillna('90%')

    inv['Lead_Time_Days'] = lead_time
    inv['Std_Dev_Daily'] = inv['Recent_Std'] / np.sqrt(30)
    inv['Safety_Stock'] = inv['Z_Score'] * inv['Std_Dev_Daily'] * np.sqrt(lead_time)
    inv['Daily_Demand'] = inv['Recent_Avg'] / 30
    inv['Min_Inventory_ROP'] = inv['Daily_Demand'] * lead_time + inv['Safety_Stock']
    inv['Order_Qty'] = inv['Recent_Avg'] * order_cycle
    inv['Max_Inventory'] = inv['Min_Inventory_ROP'] + inv['Order_Qty']
    inv['Avg_Inventory'] = (inv['Max_Inventory'] + inv['Min_Inventory_ROP']) / 2
    inv['Days_of_Supply_Min'] = np.where(inv['Daily_Demand'] > 0, inv['Min_Inventory_ROP'] / inv['Daily_Demand'], 0)
    inv['Days_of_Supply_Max'] = np.where(inv['Daily_Demand'] > 0, inv['Max_Inventory'] / inv['Daily_Demand'], 0)
    inv['Days_of_Supply_Avg'] = np.where(inv['Daily_Demand'] > 0, inv['Avg_Inventory'] / inv['Daily_Demand'], 0)

    inv['Safety_Stock_Value'] = inv['Safety_Stock'] * inv['MRP'].fillna(0)
    inv['Min_Value'] = inv['Min_Inventory_ROP'] * inv['MRP'].fillna(0)
    inv['Max_Value'] = inv['Max_Inventory'] * inv['MRP'].fillna(0)
    inv['Avg_Value'] = inv['Avg_Inventory'] * inv['MRP'].fillna(0)

    for i in range(1, len(months)):
        inv[f'Dev_{months[i]}'] = qty_df[months[i]] - qty_df[months[i - 1]]
        inv[f'PctDev_{months[i]}'] = np.where(
            qty_df[months[i - 1]] > 0,
            (qty_df[months[i]] - qty_df[months[i - 1]]) / qty_df[months[i - 1]] * 100, 0)

    x = np.arange(len(recent))
    slopes = []
    for _, row in qty_df.iterrows():
        vals = row[recent].values.astype(float)
        if np.all(np.isfinite(vals)) and not np.all(vals == 0):
            slopes.append(float(np.polyfit(x, vals, 1)[0]))
        else:
            slopes.append(0.0)
    inv['Trend_Slope'] = slopes
    inv['Trend'] = np.where(inv['Trend_Slope'] > inv['Recent_Avg'] * 0.05, 'Growing',
                   np.where(inv['Trend_Slope'] < -inv['Recent_Avg'] * 0.05, 'Declining', 'Stable'))

    if idr_df is not None and len(idr_df) > 0:
        idr_agg = idr_df.groupby('Item Name')['Projection Units'].sum().reset_index()
        idr_agg.columns = ['Item_Name', 'Projection_Units']
        inv = inv.merge(idr_agg, on='Item_Name', how='left')
        inv['Forecast_Deviation_Pct'] = np.where(
            inv['Recent_Avg'] > 0,
            (inv['Projection_Units'] - inv['Recent_Avg']) / inv['Recent_Avg'] * 100, np.nan)
    else:
        inv['Projection_Units'] = np.nan
        inv['Forecast_Deviation_Pct'] = np.nan

    inv = inv.sort_values(['ABC_Category', 'Annual_Usage_Value'], ascending=[True, False])
    return inv, months


def build_formula_excel(inv_df, months, lead_time, z_scores, order_cycle):
    wb = Workbook()
    hfill = PatternFill('solid', fgColor='1F4E79')
    hfont = Font(bold=True, color='FFFFFF', size=10, name='Arial')
    bfont = Font(size=10, name='Arial', color='0000FF')
    dfont = Font(size=10, name='Arial')

    n_months = len(months)
    n_rows = len(inv_df)
    first_m_col = 8
    last_m_col = 7 + n_months
    recent_start = last_m_col - 5
    fm = get_column_letter(first_m_col)
    lm = get_column_letter(last_m_col)
    rs = get_column_letter(recent_start)

    # Sheet 1: Raw Data
    ws = wb.active
    ws.title = 'Raw_Data'
    raw_h = ['Item_Name', 'ABC_Category', 'Item_Group', 'Origin', 'MRP', 'Z_Score', 'Lead_Time'] + months
    ws.append(raw_h)
    for cell in ws[1]:
        cell.fill = hfill
        cell.font = hfont
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    for idx, (_, row) in enumerate(inv_df.iterrows(), 2):
        ws.cell(row=idx, column=1, value=row['Item_Name']).font = dfont
        ws.cell(row=idx, column=2, value=row.get('ABC_Category', '')).font = dfont
        ws.cell(row=idx, column=3, value=row['New MIS ITEM Group']).font = dfont
        ws.cell(row=idx, column=4, value=row.get('Origin', '')).font = dfont
        ws.cell(row=idx, column=5, value=row.get('MRP', 0)).font = bfont
        ws.cell(row=idx, column=6, value=row['Z_Score']).font = bfont
        ws.cell(row=idx, column=7, value=lead_time).font = bfont
        for mi, m in enumerate(months):
            val = row.get(m, 0)
            ws.cell(row=idx, column=8 + mi, value=val if pd.notna(val) else 0).font = dfont
            ws.cell(row=idx, column=8 + mi).number_format = '#,##0'

    ws.column_dimensions['A'].width = 55
    for c in range(2, len(raw_h) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14

    # Sheet 2: Formulas
    ws2 = wb.create_sheet('Inventory_Formulas')
    calc_h = ['Item_Name', 'ABC', 'MRP', 'Z_Score', 'LT_Days',
              'Avg_All_Months', 'Recent_6M_Avg', 'Recent_6M_StdDev', 'StdDev_Daily',
              'CoV', 'Safety_Stock', 'Daily_Demand',
              'MIN_Inventory_ROP', 'Order_Qty', 'MAX_Inventory', 'Avg_Inventory',
              'DoS_MIN', 'DoS_MAX', 'DoS_AVG',
              'SS_Value_Rs', 'MIN_Value_Rs', 'MAX_Value_Rs', 'AVG_Value_Rs']
    ws2.append(calc_h)
    for cell in ws2[1]:
        cell.fill = hfill
        cell.font = hfont
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    for r in range(2, n_rows + 2):
        ws2.cell(row=r, column=1, value=f"=Raw_Data!A{r}").font = dfont
        ws2.cell(row=r, column=2, value=f"=Raw_Data!B{r}").font = dfont
        ws2.cell(row=r, column=3, value=f"=Raw_Data!E{r}").font = bfont
        ws2.cell(row=r, column=4, value=f"=Raw_Data!F{r}").font = bfont
        ws2.cell(row=r, column=5, value=f"=Raw_Data!G{r}").font = bfont
        ws2.cell(row=r, column=6, value=f"=AVERAGE(Raw_Data!{fm}{r}:{lm}{r})").font = dfont
        ws2.cell(row=r, column=7, value=f"=AVERAGE(Raw_Data!{rs}{r}:{lm}{r})").font = dfont
        ws2.cell(row=r, column=8, value=f"=STDEV(Raw_Data!{rs}{r}:{lm}{r})").font = dfont
        ws2.cell(row=r, column=9, value=f"=H{r}/SQRT(30)").font = dfont
        ws2.cell(row=r, column=10, value=f"=IF(G{r}>0,H{r}/G{r},0)").font = dfont
        ws2.cell(row=r, column=11, value=f"=D{r}*I{r}*SQRT(E{r})").font = dfont
        ws2.cell(row=r, column=12, value=f"=G{r}/30").font = dfont
        ws2.cell(row=r, column=13, value=f"=L{r}*E{r}+K{r}").font = dfont
        ws2.cell(row=r, column=14, value=f"=G{r}*{order_cycle}").font = dfont
        ws2.cell(row=r, column=15, value=f"=M{r}+N{r}").font = dfont
        ws2.cell(row=r, column=16, value=f"=(O{r}+M{r})/2").font = dfont
        ws2.cell(row=r, column=17, value=f"=IF(L{r}>0,M{r}/L{r},0)").font = dfont
        ws2.cell(row=r, column=18, value=f"=IF(L{r}>0,O{r}/L{r},0)").font = dfont
        ws2.cell(row=r, column=19, value=f"=IF(L{r}>0,P{r}/L{r},0)").font = dfont
        ws2.cell(row=r, column=20, value=f"=K{r}*C{r}").font = dfont
        ws2.cell(row=r, column=21, value=f"=M{r}*C{r}").font = dfont
        ws2.cell(row=r, column=22, value=f"=O{r}*C{r}").font = dfont
        ws2.cell(row=r, column=23, value=f"=P{r}*C{r}").font = dfont
        for c in [6, 7, 8, 9, 11, 12, 13, 14, 15, 16]:
            ws2.cell(row=r, column=c).number_format = '#,##0'
        ws2.cell(row=r, column=10).number_format = '0.000'
        for c in [17, 18, 19]:
            ws2.cell(row=r, column=c).number_format = '0.0'
        for c in [20, 21, 22, 23]:
            ws2.cell(row=r, column=c).number_format = '#,##0'

    ws2.column_dimensions['A'].width = 55
    for c in range(2, 24):
        ws2.column_dimensions[get_column_letter(c)].width = 16

    # Sheet 3: Deviations with formulas
    ws3 = wb.create_sheet('Deviations_Formulas')
    n_dev = n_months - 1
    dev_h = ['Item_Name', 'ABC'] + [f'Dev_{months[i]}' for i in range(1, n_months)] + [f'%Dev_{months[i]}' for i in range(1, n_months)]
    ws3.append(dev_h)
    for cell in ws3[1]:
        cell.fill = hfill
        cell.font = hfont
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    for r in range(2, n_rows + 2):
        ws3.cell(row=r, column=1, value=f"=Raw_Data!A{r}").font = dfont
        ws3.cell(row=r, column=2, value=f"=Raw_Data!B{r}").font = dfont
        for i in range(n_dev):
            curr = get_column_letter(first_m_col + i + 1)
            prev = get_column_letter(first_m_col + i)
            ws3.cell(row=r, column=3 + i, value=f"=Raw_Data!{curr}{r}-Raw_Data!{prev}{r}").font = dfont
            ws3.cell(row=r, column=3 + i).number_format = '#,##0'
            ws3.cell(row=r, column=3 + n_dev + i,
                     value=f"=IF(Raw_Data!{prev}{r}>0,(Raw_Data!{curr}{r}-Raw_Data!{prev}{r})/Raw_Data!{prev}{r}*100,0)").font = dfont
            ws3.cell(row=r, column=3 + n_dev + i).number_format = '0.0'

    ws3.column_dimensions['A'].width = 55

    # Sheet 4: Legend
    ws4 = wb.create_sheet('Legend')
    legend = [
        ['FARMLEY INVENTORY MODEL - FORMULA REFERENCE', '', ''],
        ['', '', ''],
        ['Column', 'Formula', 'Description'],
        ['Safety Stock (K)', '= Z x StdDev_Daily x SQRT(LT)', 'Buffer stock for demand uncertainty'],
        ['StdDev Daily (I)', '= StdDev_Monthly / SQRT(30)', 'Monthly to daily conversion'],
        ['MIN / ROP (M)', '= (Daily_Demand x LT) + Safety_Stock', 'Reorder trigger'],
        ['Order Qty (N)', f'= Recent_6M_Avg x {order_cycle}', 'Cycle replenishment quantity'],
        ['MAX (O)', '= MIN + Order_Qty', 'Post-replenishment upper bound'],
        ['Avg Inventory (P)', '= (MAX + MIN) / 2', 'Expected inventory at any time'],
        ['', '', ''],
        ['INPUT PARAMETERS (Blue = editable)', '', ''],
        ['Lead Time', f'{lead_time} days', 'Factory-warehouse co-located'],
        ['Z-Score A', f'{z_scores[0]}', f'SL: {norm_cdf(z_scores[0]):.0%}'],
        ['Z-Score B', f'{z_scores[1]}', f'SL: {norm_cdf(z_scores[1]):.0%}'],
        ['Z-Score C', f'{z_scores[2]}', f'SL: {norm_cdf(z_scores[2]):.0%}'],
        ['Order Cycle', f'{order_cycle} month(s)', 'Replenishment frequency'],
    ]
    for ri, row_data in enumerate(legend, 1):
        for ci, val in enumerate(row_data, 1):
            cell = ws4.cell(row=ri, column=ci, value=val)
            if ri == 1:
                cell.font = Font(bold=True, size=14)
            elif ri in [3, 11]:
                cell.font = Font(bold=True)

    ws4.column_dimensions['A'].width = 35
    ws4.column_dimensions['B'].width = 45
    ws4.column_dimensions['C'].width = 50

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def safe_json(val):
    if isinstance(val, float) and (np.isnan(val) or np.isinf(val)):
        return None
    if isinstance(val, (np.integer,)):
        return int(val)
    if isinstance(val, (np.floating,)):
        return float(val)
    return val


@app.post("/api/process")
@app.post("/process")
async def process_files(
    sigma_file: UploadFile = File(...),
    abc_file: UploadFile = File(...),
    idr_file: UploadFile = File(None),
    lead_time: float = Form(2),
    z_a: float = Form(1.65),
    z_b: float = Form(1.28),
    z_c: float = Form(1.04),
    order_cycle: float = Form(1.0),
):
    sigma_bytes = await sigma_file.read()
    abc_bytes = await abc_file.read()

    sigma_df = pd.read_excel(BytesIO(sigma_bytes), sheet_name=0)
    abc_df = pd.read_excel(BytesIO(abc_bytes), sheet_name='abc_summary').dropna(subset=['Item_Name'])

    idr_df = None
    if idr_file and idr_file.filename:
        idr_bytes = await idr_file.read()
        if idr_bytes:
            try:
                idr_df = pd.read_excel(BytesIO(idr_bytes), sheet_name='Projection')
            except Exception:
                idr_df = None

    z_scores = [z_a, z_b, z_c]
    inv_df, months = compute_inventory(sigma_df, abc_df, idr_df, lead_time, z_scores, order_cycle)

    CACHE['inv_df'] = inv_df
    CACHE['months'] = months
    CACHE['sigma_bytes'] = sigma_bytes
    CACHE['abc_bytes'] = abc_bytes
    CACHE['idr_bytes'] = idr_bytes if idr_file and idr_file.filename else None
    CACHE['params'] = {'lead_time': lead_time, 'z_scores': z_scores, 'order_cycle': order_cycle}

    display_cols = ['Item_Name', 'ABC_Category', 'New MIS ITEM Group', 'Origin', 'MRP',
                    'Service_Level', 'Lead_Time_Days', 'Recent_Avg', 'Recent_Std', 'CoV',
                    'Z_Score', 'Safety_Stock', 'Daily_Demand',
                    'Min_Inventory_ROP', 'Order_Qty', 'Max_Inventory', 'Avg_Inventory',
                    'Days_of_Supply_Min', 'Days_of_Supply_Max', 'Days_of_Supply_Avg',
                    'Safety_Stock_Value', 'Min_Value', 'Max_Value', 'Avg_Value', 'Trend']

    existing_cols = [c for c in display_cols if c in inv_df.columns]
    rows = []
    for _, r in inv_df[existing_cols].iterrows():
        rows.append({c: safe_json(r[c]) for c in existing_cols})

    dev_months = [m for m in months[1:]]
    dev_data = []
    for _, r in inv_df.iterrows():
        d = {'Item_Name': r['Item_Name'], 'ABC_Category': safe_json(r.get('ABC_Category', '')),
             'New MIS ITEM Group': r['New MIS ITEM Group']}
        for m in dev_months:
            d[f'Dev_{m}'] = safe_json(r.get(f'Dev_{m}', 0))
            d[f'PctDev_{m}'] = safe_json(r.get(f'PctDev_{m}', 0))
        dev_data.append(d)

    forecast_data = []
    if 'Projection_Units' in inv_df.columns:
        fc = inv_df[inv_df['Projection_Units'].notna()]
        for _, r in fc.iterrows():
            dev_pct = safe_json(r.get('Forecast_Deviation_Pct', None))
            risk = 'N/A'
            if dev_pct is not None:
                risk = 'HIGH' if abs(dev_pct) > 50 else ('MEDIUM' if abs(dev_pct) > 25 else 'LOW')
            forecast_data.append({
                'Item_Name': r['Item_Name'],
                'ABC_Category': safe_json(r.get('ABC_Category', '')),
                'New MIS ITEM Group': r['New MIS ITEM Group'],
                'Recent_Avg': safe_json(r['Recent_Avg']),
                'Projection_Units': safe_json(r['Projection_Units']),
                'Deviation_Units': safe_json(r['Projection_Units'] - r['Recent_Avg']),
                'Deviation_Pct': dev_pct,
                'Risk': risk,
            })

    groups = sorted(inv_df['New MIS ITEM Group'].dropna().unique().tolist())
    origins = sorted(inv_df['Origin'].dropna().unique().tolist())

    summary = {}
    for cat in ['A', 'B', 'C']:
        s = inv_df[inv_df['ABC_Category'] == cat]
        if len(s) > 0:
            summary[cat] = {
                'count': len(s),
                'sl': {'A': f"{norm_cdf(z_a):.0%}", 'B': f"{norm_cdf(z_b):.0%}", 'C': f"{norm_cdf(z_c):.0%}"}[cat],
                'avg_ss': round(float(s['Safety_Stock'].mean()), 0),
                'total_ss_cr': round(float(s['Safety_Stock_Value'].sum()) / 1e7, 2),
                'total_min_cr': round(float(s['Min_Value'].sum()) / 1e7, 2),
                'total_max_cr': round(float(s['Max_Value'].sum()) / 1e7, 2),
                'avg_dos': round(float(s['Days_of_Supply_Avg'].mean()), 1),
                'growing': int((s['Trend'] == 'Growing').sum()),
                'declining': int((s['Trend'] == 'Declining').sum()),
                'stable': int((s['Trend'] == 'Stable').sum()),
            }

    totals = {
        'skus': len(inv_df),
        'ss_value_cr': round(float(inv_df['Safety_Stock_Value'].sum()) / 1e7, 2),
        'min_value_cr': round(float(inv_df['Min_Value'].sum()) / 1e7, 2),
        'max_value_cr': round(float(inv_df['Max_Value'].sum()) / 1e7, 2),
        'avg_value_cr': round(float(inv_df['Avg_Value'].sum()) / 1e7, 2),
    }

    return JSONResponse({
        'inventory': rows,
        'deviations': dev_data,
        'forecast': forecast_data,
        'summary': summary,
        'totals': totals,
        'filters': {'groups': groups, 'origins': origins},
        'months': months,
        'dev_months': dev_months,
    })


@app.get("/api/download")
@app.get("/download")
async def download_excel():
    if 'inv_df' not in CACHE:
        return JSONResponse({'error': 'No data processed yet'}, status_code=400)

    inv_df = CACHE['inv_df']
    months = CACHE['months']
    p = CACHE['params']

    buf = build_formula_excel(inv_df, months, p['lead_time'], p['z_scores'], p['order_cycle'])

    return StreamingResponse(
        buf,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.xml',
        headers={'Content-Disposition': 'attachment; filename=Farmley_Inventory_Model.xlsx'}
    )
